"""Edits a document in place, keeping its own styles.

    doc_edit.py <path> --ops '<json array>' [--out <path>]

.docx edits go through python-docx, so the template, numbering, headers and
every style the document already has stay as they were: a replaced phrase
keeps the formatting of the run it started in, an inserted paragraph takes a
named style from the document itself. .doc/.odt/.rtf are converted to .docx
first and saved as .docx.

A PDF's body text cannot be edited in place -- it is positioned glyphs, not
paragraphs. What can be done to a PDF is done here (forms, pages,
metadata); a content change means rebuilding the PDF from its source.

Every op is an object with an "op" field; unused fields may be null.
"""

from __future__ import annotations

import argparse
import copy
import json
import tempfile
from pathlib import Path

from _common import cli, emit, existing, fail, kind_of, parse_pages, soffice_convert

DOCX_OPS = {
    "replace_text", "insert_paragraph", "insert_heading", "delete_paragraph", "insert_table",
    "set_cell", "add_row", "replace_image", "insert_image", "set_properties",
}
PDF_OPS = {"fill_form", "merge", "extract_pages", "delete_pages", "rotate", "set_properties"}


# -- docx --------------------------------------------------------------------


def all_paragraphs(doc):
    """Body, table cells, headers and footers: everywhere text can be replaced."""
    yield from doc.paragraphs
    for table in doc.tables:
        yield from table_paragraphs(table)
    for section in doc.sections:
        for part in (section.header, section.footer, section.first_page_header, section.first_page_footer):
            if part is not None and not part.is_linked_to_previous:
                yield from part.paragraphs
                for table in part.tables:
                    yield from table_paragraphs(table)


def table_paragraphs(table):
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for inner in cell.tables:
                yield from table_paragraphs(inner)


def replace_in_paragraph(paragraph, find: str, text: str, every: bool) -> int:
    """Replaces across run boundaries, keeping the first matched run's formatting."""
    count = 0
    while True:
        runs = paragraph.runs
        joined = "".join(r.text for r in runs)
        at = joined.find(find)
        if at < 0 or not runs:
            return count
        end = at + len(find)
        pos = 0
        first = None
        for run in runs:
            start, stop = pos, pos + len(run.text)
            pos = stop
            if stop <= at or start >= end:
                continue
            keep_head = run.text[: max(at - start, 0)]
            keep_tail = run.text[max(end - start, 0):] if stop > end else ""
            if first is None:
                first = run
                run.text = keep_head + text + keep_tail
            else:
                run.text = keep_tail
        count += 1
        if not every:
            return count


def find_paragraph(doc, anchor: str):
    """'p:12' is the 12th body paragraph as doc_read numbers them; otherwise the first containing the text."""
    if anchor is None:
        fail("This op needs an anchor: 'p:<index>' from doc_read, or text the paragraph contains.")
    if anchor.startswith("p:"):
        index = int(anchor[2:])
        if not 0 <= index < len(doc.paragraphs):
            fail(f"Paragraph {index} does not exist; the document has {len(doc.paragraphs)}.")
        return doc.paragraphs[index]
    for paragraph in doc.paragraphs:
        if anchor in paragraph.text:
            return paragraph
    fail(f"No paragraph contains {anchor!r}. Use doc_read to find the exact text or the p:<index>.")


def styled(doc, name: str | None, fallback: str) -> str | None:
    names = {s.name for s in doc.styles}
    for candidate in (name, fallback):
        if candidate and candidate in names:
            return candidate
    return None


def place(anchor_paragraph, element, position: str | None) -> None:
    if position == "before":
        anchor_paragraph._p.addprevious(element)
    else:
        anchor_paragraph._p.addnext(element)


def new_paragraph(doc, text: str, style: str | None):
    # Built on the body and then moved, so python-docx applies the style.
    paragraph = doc.add_paragraph(text, style=style)
    return paragraph


def new_table(doc, rows: int, cols: int):
    """doc.add_table, without its assumption that the section names a page size.

    Documents written by pandoc and some generators leave the size out, and
    python-docx then fails subtracting None from None.
    """
    from docx.oxml.table import CT_Tbl
    from docx.shared import Inches
    from docx.table import Table

    section = doc.sections[-1] if doc.sections else None
    try:
        width = section.page_width - section.left_margin - section.right_margin
    except (AttributeError, TypeError):
        width = Inches(6.3)
    tbl = CT_Tbl.new_tbl(rows, cols, width)
    doc.element.body.append(tbl)
    return Table(tbl, doc._body)


def docx_table(doc, index):
    if index is None or not 1 <= index <= len(doc.tables):
        fail(f"Table {index} does not exist; the document has {len(doc.tables)} (numbered from 1).")
    return doc.tables[index - 1]


def apply_docx(doc, op: dict, workspace: Path) -> str:
    kind = op["op"]
    if kind == "replace_text":
        find, text = op.get("find"), op.get("text") or ""
        if not find:
            fail("replace_text needs 'find'.")
        n = sum(replace_in_paragraph(p, find, text, bool(op.get("all"))) for p in all_paragraphs(doc))
        if n == 0:
            fail(f"replace_text: {find!r} was not found. Text split by a field or a line break cannot match; read the document again.")
        return f"replaced {n} occurrence(s) of {find!r}"

    if kind in ("insert_paragraph", "insert_heading"):
        anchor = find_paragraph(doc, op.get("anchor"))
        if kind == "insert_heading":
            level = int(op.get("level") or 1)
            style = styled(doc, op.get("style"), f"Heading {level}")
        else:
            style = styled(doc, op.get("style"), "Normal")
        paragraph = new_paragraph(doc, op.get("text") or "", style)
        place(anchor, paragraph._p, op.get("position"))
        return f"inserted {kind.split('_')[1]} {op.get('position') or 'after'} {op.get('anchor')!r}"

    if kind == "delete_paragraph":
        paragraph = find_paragraph(doc, op.get("anchor"))
        paragraph._p.getparent().remove(paragraph._p)
        return f"deleted paragraph {op.get('anchor')!r}"

    if kind == "insert_table":
        rows = op.get("rows") or []
        if not rows:
            fail("insert_table needs 'rows', the first being the header.")
        anchor = find_paragraph(doc, op.get("anchor"))
        width = max(len(r) for r in rows)
        table = new_table(doc, len(rows), width)
        style = styled(doc, op.get("style"), "Table Grid")
        if style:
            table.style = style
        for r, values in enumerate(rows):
            for c, value in enumerate(values):
                table.cell(r, c).text = str(value)
        place(anchor, table._tbl, op.get("position"))
        return f"inserted a {len(rows)}x{width} table"

    if kind == "set_cell":
        table = docx_table(doc, op.get("table"))
        r, c = int(op.get("row") or 0), int(op.get("col") or 0)
        try:
            cell = table.cell(r - 1, c - 1)
        except IndexError:
            fail(f"Cell ({r}, {c}) is outside table {op.get('table')} ({len(table.rows)} rows, {len(table.columns)} columns; numbered from 1).")
        paragraphs = cell.paragraphs
        # Keep the first run's formatting rather than resetting the cell.
        first = paragraphs[0]
        if first.runs:
            first.runs[0].text = op.get("text") or ""
            for run in first.runs[1:]:
                run.text = ""
        else:
            first.add_run(op.get("text") or "")
        for extra in paragraphs[1:]:
            extra._p.getparent().remove(extra._p)
        return f"set table {op.get('table')} cell ({r}, {c})"

    if kind == "add_row":
        table = docx_table(doc, op.get("table"))
        values = (op.get("rows") or [[]])[0]
        # A copy of the last row keeps its borders, shading and fonts.
        last = table.rows[-1]._tr
        row = copy.deepcopy(last)
        last.addnext(row)
        new = table.rows[-1]
        for i, cell in enumerate(new.cells):
            value = str(values[i]) if i < len(values) else ""
            first = cell.paragraphs[0]
            if first.runs:
                first.runs[0].text = value
                for run in first.runs[1:]:
                    run.text = ""
            else:
                first.add_run(value)
            for extra in cell.paragraphs[1:]:
                extra._p.getparent().remove(extra._p)
        return f"added a row to table {op.get('table')}"

    if kind == "replace_image":
        index = int(op.get("index") or 1)
        image = workspace_file(workspace, op.get("image"))
        blips = doc.element.body.xpath(".//*[local-name()='blip']")
        if not 1 <= index <= len(blips):
            fail(f"Image {index} does not exist; the document body has {len(blips)}.")
        rid = blips[index - 1].get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
        part = doc.part.related_parts[rid]
        part._blob = image.read_bytes()
        return f"replaced image {index} with {image.name}"

    if kind == "insert_image":
        from docx.shared import Cm

        anchor = find_paragraph(doc, op.get("anchor"))
        image = workspace_file(workspace, op.get("image"))
        paragraph = doc.add_paragraph()
        width = op.get("width_cm")
        paragraph.add_run().add_picture(str(image), width=Cm(float(width)) if width else Cm(15))
        place(anchor, paragraph._p, op.get("position"))
        if op.get("text"):
            caption = new_paragraph(doc, op["text"], styled(doc, "Caption", "Normal"))
            paragraph._p.addnext(caption._p)
        return f"inserted {image.name}"

    if kind == "set_properties":
        props = doc.core_properties
        for field in ("title", "author", "subject"):
            if op.get(field) is not None:
                setattr(props, field, op[field])
        return "set document properties"

    fail(f"{kind} is not a .docx operation. Available: {', '.join(sorted(DOCX_OPS))}.")


def workspace_file(workspace: Path, path: str | None) -> Path:
    if not path:
        fail("This op needs 'image', a path in the workspace.")
    resolved = (workspace / path).resolve()
    if workspace not in resolved.parents:
        fail(f"{path} is outside the workspace.")
    return existing(str(resolved))


# -- pdf ---------------------------------------------------------------------


def apply_pdf(path: Path, ops: list[dict], out: Path, workspace: Path) -> list[str]:
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter(clone_from=str(path))
    done: list[str] = []
    for op in ops:
        kind = op["op"]
        if kind == "fill_form":
            fields = {f["name"]: f["value"] for f in (op.get("fields") or [])}
            if not fields:
                fail("fill_form needs 'fields': [{name, value}]. doc_read lists the form's field names.")
            for page in writer.pages:
                writer.update_page_form_field_values(page, fields, auto_regenerate=False)
            writer.set_need_appearances_writer(True)
            done.append(f"filled {len(fields)} field(s)")
        elif kind == "merge":
            for name in op.get("files") or []:
                writer.append(str(workspace_file(workspace, name)))
            done.append(f"appended {len(op.get('files') or [])} file(s)")
        elif kind in ("extract_pages", "delete_pages"):
            pages = parse_pages(op.get("pages"), len(writer.pages))
            if not pages:
                fail(f"{kind} needs 'pages', e.g. '1-3,5'.")
            keep = pages if kind == "extract_pages" else [p for p in range(1, len(writer.pages) + 1) if p not in pages]
            fresh = PdfWriter()
            for number in keep:
                fresh.add_page(writer.pages[number - 1])
            writer = fresh
            done.append(f"{'kept' if kind == 'extract_pages' else 'deleted'} pages {op.get('pages')}")
        elif kind == "rotate":
            degrees = int(op.get("degrees") or 90)
            for number in parse_pages(op.get("pages"), len(writer.pages)):
                writer.pages[number - 1].rotate(degrees)
            done.append(f"rotated pages {op.get('pages') or 'all'} by {degrees}")
        elif kind == "set_properties":
            meta = {f"/{k.capitalize()}": op[k] for k in ("title", "author", "subject") if op.get(k) is not None}
            writer.add_metadata(meta)
            done.append("set document properties")
        else:
            hint = " A PDF's text cannot be edited in place: rebuild it from its Markdown source with doc_render, or edit the .docx it came from." if kind in DOCX_OPS else ""
            fail(f"{kind} is not a PDF operation. Available: {', '.join(sorted(PDF_OPS))}.{hint}")
    with open(out, "wb") as handle:
        writer.write(handle)
    return done


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("path")
    parser.add_argument("--ops", required=True)
    parser.add_argument("--out")
    parser.add_argument("--workspace", default=".")
    args = parser.parse_args()

    try:
        ops = json.loads(args.ops)
    except json.JSONDecodeError as error:
        fail(f"--ops is not valid JSON: {error}")
    if not isinstance(ops, list) or not ops or not all(isinstance(o, dict) and o.get("op") for o in ops):
        fail("--ops must be a non-empty array of objects, each with an 'op'.")

    path = existing(args.path)
    workspace = Path(args.workspace).resolve()
    kind = kind_of(path)
    out = Path(args.out) if args.out else path

    if kind == "pdf":
        done = apply_pdf(path, ops, out, workspace)
        emit({"path": str(out), "applied": done})

    if kind == "text":
        fail(f"{path.name} is plain text; edit it with edit_file.")

    from docx import Document

    with tempfile.TemporaryDirectory() as scratch:
        source = path
        if kind != "docx":
            source = soffice_convert(path, "docx", Path(scratch))
            if not args.out:
                out = path.with_suffix(".docx")
        doc = Document(str(source))
        done = [apply_docx(doc, op, workspace) for op in ops]
        out.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(out))
    emit({"path": str(out), "applied": done})


if __name__ == "__main__":
    cli(main)
